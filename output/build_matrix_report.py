from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path('/Users/osmansamilerdogan/Desktop/midtermexam')
OUTPUT_PATH = WORKSPACE / 'Direct_Manipulation_Matrix_Sorting_Report.docx'
INITIAL_SCREENSHOT = WORKSPACE / 'output/playwright/matrix-initial.png'
GESTURE_SCREENSHOT = WORKSPACE / 'output/playwright/matrix-diagonal-gesture.png'
SORTED_SCREENSHOT = WORKSPACE / 'output/playwright/matrix-after-sort.png'

NAVY = '17365D'
BLUE = '2563EB'
PALE_BLUE = 'EAF2F8'
PALE_GRAY = 'F5F7FA'
MID_GRAY = 'D9D9D9'
DARK_GRAY = '3F4A56'
WHITE = 'FFFFFF'
BLACK = '000000'


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        tc_pr.append(shading)
    shading.set(qn('w:fill'), color)


def set_cell_borders(cell, color=MID_GRAY, size='6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        item = borders.find(tag)
        if item is None:
            item = OxmlElement(f'w:{edge}')
            borders.append(item)
        item.set(qn('w:val'), 'single')
        item.set(qn('w:sz'), size)
        item.set(qn('w:color'), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, amount in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(amount))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, keep=True):
    p_pr = paragraph._p.get_or_add_pPr()
    item = p_pr.find(qn('w:keepNext'))
    if keep and item is None:
        p_pr.append(OxmlElement('w:keepNext'))
    elif not keep and item is not None:
        p_pr.remove(item)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True,
    )
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), relationship_id)
    run = OxmlElement('w:r')
    run_properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), BLUE)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_element = OxmlElement('w:t')
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Page ')
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    number = OxmlElement('w:t')
    number.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, separate, number, end])


def add_title(doc, title, subtitle=None):
    paragraph = doc.add_paragraph(style='Title')
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(16)
    run = paragraph.add_run(title)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.paragraph_format.space_after = Pt(14)
        sub_run = sub.add_run(subtitle)
        sub_run.font.size = Pt(16)
        sub_run.font.color.rgb = RGBColor(63, 74, 86)


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph(style='Body Text')
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(f'{index}.').bold = True
        p.add_run(f'  {item}')


def add_code(doc, code):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), PALE_GRAY)
    p_pr.append(shading)
    for index, line in enumerate(code.strip('\n').split('\n')):
        run = paragraph.add_run(line)
        run.font.name = 'Menlo'
        run._element.rPr.rFonts.set(qn('w:ascii'), 'Menlo')
        run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Menlo')
        run.font.size = Pt(8.5)
        if index < len(code.strip('\n').split('\n')) - 1:
            run.add_break()
    return paragraph


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(63, 74, 86)
    return p


def add_figure(doc, path, caption, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_keep_with_next(p)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(6.8))
    c_nv_pr = shape._inline.docPr
    c_nv_pr.set('descr', alt_text)
    add_caption(doc, caption)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].height_rule = None
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        if widths:
            cell.width = widths[index]
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate(row_values):
            cell = row.cells[col_index]
            if widths:
                cell.width = widths[col_index]
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            run.font.size = Pt(8.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.76)
    section.right_margin = Inches(0.76)
    section.different_first_page_header_footer = True

    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    body = doc.styles['Body Text']
    body.font.name = 'Aptos'
    body._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    body._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    body.font.size = Pt(10.5)
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.08

    title_style = doc.styles['Title']
    title_style.font.name = 'Aptos Display'
    title_style._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
    title_style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
    title_style.font.size = Pt(30)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn('w:pBdr'))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (('Heading 1', 17), ('Heading 2', 13)):
        style = doc.styles[style_name]
        style.font.name = 'Aptos Display'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if style_name == 'Heading 1' else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    for list_style in ('List Bullet', 'List Number'):
        style = doc.styles[list_style]
        style.font.name = 'Aptos'
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = 'Direct Manipulation Matrix Sorting Application'
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.runs[0].font.size = Pt(8.5)
    header_p.runs[0].font.color.rgb = RGBColor(90, 100, 110)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    add_page_number(footer_p)
    for run in footer_p.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(90, 100, 110)

    doc.core_properties.title = 'Direct Manipulation Matrix Sorting Application'
    doc.core_properties.subject = 'Implementation process and design report for a 30 by 40 sortable matrix'
    doc.core_properties.keywords = 'direct manipulation, matrix sorting, React, Vite, pointer events'

    doc.add_paragraph().paragraph_format.space_after = Pt(36)
    add_title(doc, 'Direct Manipulation Matrix Sorting Application', 'Implementation Process and Design Report')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run('CSCI 626 Midterm Examination')
    r.bold = True
    r.font.size = Pt(13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('React  Vite  Tailwind CSS  SVG').font.size = Pt(11)

    doc.add_paragraph().paragraph_format.space_after = Pt(32)
    add_body(
        doc,
        'This report explains how the submitted web application was created and how its drag gestures sort a 30 by 40 numerical matrix. The implementation displays all 1,200 values at once and performs row, column, and combined sorting without buttons, menus, sliders, checkboxes, tooltips, or other interface widgets.',
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    run = p.add_run('Prepared from the verified project source and production build')
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(63, 74, 86)
    doc.add_page_break()

    heading(doc, 'Report Contents')
    add_numbered(doc, [
        'Project objective and requirement compliance',
        'Technology choice and project creation',
        'Application structure and matrix data',
        'Direct manipulation and gesture classification',
        'Sorting algorithms and visual transitions',
        'User instructions and test procedure',
        'Engineering practices and references',
    ])

    heading(doc, 'Project Objective')
    add_body(
        doc,
        'The application turns the matrix itself into the control surface. A user presses a cell and drags in the direction that expresses the required operation. The starting cell supplies the row and column used as sort keys. This keeps the interaction inside the required matrix area and removes the need for conventional widgets.',
    )
    add_body(
        doc,
        'The values remain attached to their original rows and columns. Sorting changes the displayed order through rowOrder and colOrder arrays; it does not overwrite the underlying 30 by 40 data array. This separation makes the behavior easier to understand, animate, and test.',
    )

    heading(doc, 'Requirement Compliance')
    add_table(
        doc,
        ['Requirement', 'Implementation'],
        [
            ['30 rows and 40 columns', 'Constants ROWS = 30 and COLS = 40 create exactly 1,200 cells.'],
            ['Numerical values', 'Every cell receives an integer from 0 through 99.'],
            ['Sort rows', 'A vertical drag reorders complete rows using the starting column as the key.'],
            ['Sort columns', 'A horizontal drag reorders complete columns using the starting row as the key.'],
            ['Sort both axes', 'A diagonal drag performs the row sort and column sort from the same starting cell.'],
            ['Ascending and descending', 'Down and right mean ascending; up and left mean descending.'],
            ['No UI widgets', 'The page contains instructions, the matrix, and an SVG feedback arrow only.'],
            ['Instructions on the page', 'A short explanation above the matrix describes every gesture and direction.'],
        ],
        [Inches(2.05), Inches(4.65)],
    )

    doc.add_page_break()
    heading(doc, 'Technology Choice')
    add_body(
        doc,
        'The prompt gives D3.js as an example rather than a required library. This implementation uses React for component state and rendering, Vite for development and production builds, Tailwind CSS for utility styling, and inline SVG for the drag arrow. These browser technologies are sufficient because the visualization needs direct cell positioning and pointer input rather than scales, axes, or chart layouts.',
    )
    add_bullets(doc, [
        'React stores the row order, column order, selected cell, and drag line as component state.',
        'Vite provides the development server and creates the optimized production bundle.',
        'Tailwind CSS supplies concise layout, typography, color, and transition classes.',
        'Pointer Events provide one interaction model for mouse, pen, and touch input.',
        'SVG draws a direction arrow without adding a button or another widget.',
    ])

    heading(doc, 'Project Creation')
    add_body(doc, 'The following sequence recreates the project from an empty working directory.')
    add_numbered(doc, [
        'Create a Vite project with the React template and enter the new matrix-sort directory.',
        'Install React, React DOM, Vite, the React plugin, Tailwind CSS, and the Tailwind Vite plugin.',
        'Configure Vite to load both the React and Tailwind plugins.',
        'Import Tailwind CSS from src/index.css.',
        'Create the React root in src/main.jsx and render the App component.',
        'Implement the matrix, gesture recognition, sorting logic, and instructions in src/App.jsx.',
        'Run the development server during implementation and create a production build before submission.',
    ])
    add_code(doc, """npm create vite@latest matrix-sort -- --template react
cd matrix-sort
npm install
npm install tailwindcss @tailwindcss/vite
npm run dev
npm run build""")
    add_body(
        doc,
        'The submitted package.json already defines dev, build, and preview scripts. The verified build completed successfully with 31 transformed modules and generated the production files in the dist directory.',
    )

    heading(doc, 'Project Structure')
    add_table(
        doc,
        ['File', 'Purpose'],
        [
            ['index.html', 'Defines the browser document, the root element, and the JavaScript module entry point.'],
            ['src/main.jsx', 'Creates the React root, imports the stylesheet, and renders App.'],
            ['src/App.jsx', 'Contains the matrix model, gesture handlers, sorting operations, SVG feedback, and cell rendering.'],
            ['src/index.css', 'Imports Tailwind CSS.'],
            ['vite.config.js', 'Registers the React and Tailwind Vite plugins.'],
            ['package.json', 'Records dependencies and development, build, and preview commands.'],
        ],
        [Inches(1.55), Inches(5.15)],
    )

    heading(doc, 'Matrix Data Model')
    add_body(
        doc,
        'The application defines the dimensions and cell geometry once. The range helper creates integer index sequences. makeMatrix then creates 30 arrays, each containing 40 random integers. useMemo calls makeMatrix only for the initial component instance, so normal interaction does not replace the values with a new random matrix.',
    )
    add_code(doc, """const ROWS = 30;
const COLS = 40;

const range = (n) => Array.from({ length: n }, (_, i) => i);

function makeMatrix() {
  return range(ROWS).map(() =>
    range(COLS).map(() => Math.floor(Math.random() * 100))
  );
}""")
    add_body(
        doc,
        'The data array stores the values. Two separate state arrays store the current display order. rowOrder begins as 0 through 29, and colOrder begins as 0 through 39. Sorting copies and reorders these index arrays, which preserves the original values and row or column membership.',
    )
    add_code(doc, """const data = useMemo(makeMatrix, []);
const [rowOrder, setRowOrder] = useState(range(ROWS));
const [colOrder, setColOrder] = useState(range(COLS));""")

    heading(doc, 'Cell Rendering and Visual Encoding')
    add_body(
        doc,
        'The component maps across the original data array and creates 1,200 absolutely positioned div elements. The rowPos and colPos lookup objects translate original row and column indexes into their current screen positions. CSS transforms move each cell to its ordered position. The transition class animates a sort for 500 milliseconds so the user can see the rows and columns move.',
    )
    add_body(
        doc,
        'Blue lightness encodes cell magnitude: larger values appear darker. During a drag, the selected cell turns yellow and the complete key row and key column use a yellow scale. White cell borders preserve the grid structure, and high values use white text for contrast.',
    )
    add_figure(
        doc,
        INITIAL_SCREENSHOT,
        'Figure 1  Initial application view with instructions and all 1,200 matrix values',
        'Initial matrix interface showing written instructions and a 30 row by 40 column matrix.',
    )

    doc.add_page_break()
    heading(doc, 'Direct Manipulation Design')
    add_body(
        doc,
        'The user performs each command directly on a matrix cell. The gesture begins with the sort key and ends with the desired direction. Vertical motion changes row order, horizontal motion changes column order, and diagonal motion changes both. This design follows the direct manipulation principle of acting on visible objects with immediate visual feedback.',
    )
    add_table(
        doc,
        ['Gesture', 'Sort operation', 'Direction result', 'Feedback color'],
        [
            ['Drag down', 'Rows by starting column', 'Ascending', 'Blue'],
            ['Drag up', 'Rows by starting column', 'Descending', 'Blue'],
            ['Drag right', 'Columns by starting row', 'Ascending', 'Green'],
            ['Drag left', 'Columns by starting row', 'Descending', 'Green'],
            ['Drag diagonally', 'Rows and columns', 'Each axis follows its drag sign', 'Purple'],
        ],
        [Inches(1.35), Inches(2.15), Inches(1.85), Inches(1.35)],
    )

    heading(doc, 'Gesture Classification')
    add_body(
        doc,
        'classifyGesture compares the absolute horizontal and vertical distances. A drag shorter than 15 pixels is ignored, which prevents an ordinary click or small hand movement from causing a sort. A gesture is diagonal when its smaller axis is at least half the length of its larger axis. All other gestures are classified according to their dominant axis.',
    )
    add_code(doc, """function classifyGesture(dx, dy) {
  const ax = Math.abs(dx);
  const ay = Math.abs(dy);
  if (Math.max(ax, ay) < MIN_DRAG) return null;
  const isDiagonalGesture =
    Math.min(ax, ay) / Math.max(ax, ay) >= DIAGONAL_RATIO;
  if (isDiagonalGesture) return 'diagonal';
  return ay > ax ? 'vertical' : 'horizontal';
}""")
    add_body(
        doc,
        'The arrow uses the same classification function while the pointer is moving. Its color therefore previews the operation that will occur when the pointer is released.',
    )
    add_figure(
        doc,
        GESTURE_SCREENSHOT,
        'Figure 2  Diagonal drag with a purple arrow and highlighted key row and column',
        'Matrix during a diagonal drag showing a purple arrow and yellow selected row and column.',
    )

    doc.add_page_break()
    heading(doc, 'Pointer Event Sequence')
    add_body(
        doc,
        'The matrix container handles the complete pointer lifecycle. This reduces repeated event listeners and keeps the gesture state in one component.',
    )
    add_numbered(doc, [
        'onPointerDown reads the data-row and data-col attributes from the pressed cell.',
        'The handler converts the pointer location into coordinates relative to the matrix container.',
        'setPointerCapture keeps subsequent pointer events attached to the container while the drag continues.',
        'onPointerMove updates the SVG line and causes the row and column highlights to remain visible.',
        'onPointerUp calculates dx and dy, classifies the gesture, clears the temporary feedback, and runs the requested sort.',
        'onPointerCancel clears the same temporary state if the browser cancels the interaction.',
    ])
    add_code(doc, """if (gesture === 'vertical' || gesture === 'diagonal')
  sortRows(s.col, dy > 0);

if (gesture === 'horizontal' || gesture === 'diagonal')
  sortCols(s.row, dx > 0);""")
    add_body(
        doc,
        'The signs of dy and dx encode ascending or descending order. A positive vertical displacement means down and therefore ascending rows. A positive horizontal displacement means right and therefore ascending columns.',
    )

    heading(doc, 'Row Sorting')
    add_body(
        doc,
        'sortRows receives the original column index of the pressed cell and a Boolean direction. It copies the current row order, compares the values from the chosen column, and updates the row order state. The multiplier is 1 for ascending order and negative 1 for descending order.',
    )
    add_code(doc, """const sortRows = (keyCol, asc) =>
  setRowOrder((order) =>
    [...order].sort((a, b) =>
      (asc ? 1 : -1) * (data[a][keyCol] - data[b][keyCol])
    )
  );""")

    heading(doc, 'Column Sorting')
    add_body(
        doc,
        'sortCols follows the same pattern but compares column indexes by reading values from the original row of the pressed cell. Only colOrder changes, so complete columns move together.',
    )
    add_code(doc, """const sortCols = (keyRow, asc) =>
  setColOrder((order) =>
    [...order].sort((a, b) =>
      (asc ? 1 : -1) * (data[keyRow][a] - data[keyRow][b])
    )
  );""")

    heading(doc, 'Combined Row and Column Sorting')
    add_body(
        doc,
        'A diagonal release calls both functions. The starting row and column are saved before either sort begins, so the first state update cannot change the key used by the second operation. React applies the two state changes and renders the new row and column positions together.',
    )

    doc.add_page_break()
    heading(doc, 'Visual Reordering')
    add_body(
        doc,
        'The rowPos and colPos objects invert the order arrays. Each original index maps to its current visible position. The transform expression multiplies those positions by the fixed cell width and height. React keeps each cell identity stable with a key made from its original row and column, so the browser animates existing cells instead of replacing them.',
    )
    add_code(doc, """const rowPos = useMemo(
  () => Object.fromEntries(rowOrder.map((r, i) => [r, i])),
  [rowOrder]
);

transform: `translate(${colPos[c] * CELL_W}px,
                      ${rowPos[r] * CELL_H}px)`""")
    add_figure(
        doc,
        SORTED_SCREENSHOT,
        'Figure 3  Matrix after the diagonal sort has reordered both axes',
        'Matrix after a diagonal sorting gesture showing changed row and column positions.',
    )

    heading(doc, 'Instructions Presented to the User')
    add_body(
        doc,
        'The page tells the user to press a cell, drag, and release. It explains that the selected row and column provide the sort keys and that the values inside those groups are not independently rearranged. The three gesture descriptions use the same blue, green, and purple colors as the feedback arrow. This creates a visible connection between the written instructions and the live gesture preview.',
    )
    add_body(
        doc,
        'Because the assignment penalizes every displayed widget, the design contains no button, select menu, slider, checkbox, tooltip, text field, legend control, or reset control. The text and matrix provide all required information and interaction.',
    )

    heading(doc, 'Testing and Verification')
    add_body(doc, 'The implementation can be checked with the following procedure.')
    add_numbered(doc, [
        'Run npm install from the matrix-sort directory if dependencies are not already installed.',
        'Run npm run dev and open the local address shown by Vite.',
        'Confirm that the matrix contains 30 visible rows and 40 visible columns.',
        'Drag downward from any cell and confirm that the values in its original column become ascending from top to bottom.',
        'Drag upward and confirm that the same row sort is descending.',
        'Drag right from any cell and confirm that the values in its original row become ascending from left to right.',
        'Drag left and confirm that the same column sort is descending.',
        'Drag diagonally in each direction and confirm that both axes change according to the horizontal and vertical signs.',
        'Make a movement shorter than 15 pixels and confirm that no sorting occurs.',
        'Run npm run build and confirm that Vite produces the dist directory without errors.',
    ])
    add_body(
        doc,
        'Verification performed for this report confirmed a successful production build. A browser snapshot also exposed all 1,200 numerical cell nodes, and a diagonal pointer gesture produced the expected purple arrow, key-line highlighting, and final matrix reordering. The only browser console error was a missing optional favicon, which does not affect matrix behavior.',
    )

    heading(doc, 'Software Engineering Practices')
    add_bullets(doc, [
        'Constants define dimensions, geometry, thresholds, and gesture colors in one location.',
        'Functions have focused responsibilities for data generation, gesture classification, row sorting, column sorting, pointer conversion, and drag cleanup.',
        'State names describe the information they hold: rowOrder, colOrder, selected, and line.',
        'The code copies arrays before sorting, which avoids mutating React state.',
        'A targeted comment explains why diagonal sorting must retain the originally pressed row and column.',
        'Pointer cancellation and minimum drag handling prevent unintended actions.',
        'Stable cell keys and transform transitions preserve visual continuity during sorting.',
    ])

    heading(doc, 'Limitations and Possible Improvements')
    add_body(
        doc,
        'The values are random on each full page load, so a screenshot from one session cannot be compared numerically with a later reload. Equal values depend on the JavaScript engine’s stable array sort behavior. The fixed 26 by 20 pixel cells also assume a desktop-width viewport; a narrow display may require horizontal scrolling. These constraints do not prevent the required sorting operations.',
    )
    add_body(
        doc,
        'A later version could add keyboard gestures or screen-reader instructions while preserving the prohibition on visible widgets. Any such extension should retain direct manipulation as the primary interaction and should not add controls that violate the assignment.',
    )

    doc.add_page_break()
    heading(doc, 'References')
    references = [
        ('Shneiderman, B. 1983. Direct Manipulation A Step Beyond Programming Languages. Computer 16 8 57 to 69. DOI 10.1109 MC 1983 1654471.', 'https://doi.org/10.1109/MC.1983.1654471'),
        ('React documentation for useState.', 'https://react.dev/reference/react/useState'),
        ('React documentation for useMemo.', 'https://react.dev/reference/react/useMemo'),
        ('React documentation for useRef.', 'https://react.dev/reference/react/useRef'),
        ('Vite getting started guide.', 'https://vite.dev/guide/'),
        ('Vite production build guide.', 'https://vite.dev/guide/build'),
        ('Tailwind CSS installation with Vite.', 'https://tailwindcss.com/docs/installation/using-vite'),
        ('MDN documentation for Element setPointerCapture.', 'https://developer.mozilla.org/en-US/docs/Web/API/Element/setPointerCapture'),
    ]
    for index, (label, url) in enumerate(references, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(6)
        p.add_run(f'{index}. {label} ')
        add_hyperlink(p, 'Open source', url)

    heading(doc, 'Source Note')
    add_body(
        doc,
        'The gesture mapping and application code are implemented in the submitted project rather than copied from a third-party example. The references above document the direct manipulation concept and the framework and browser APIs used to build the application.',
    )

    doc.save(OUTPUT_PATH)


if __name__ == '__main__':
    build_document()
    print(OUTPUT_PATH)
